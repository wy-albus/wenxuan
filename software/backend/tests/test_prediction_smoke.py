from __future__ import annotations

from pathlib import Path
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from fastapi.testclient import TestClient


def _monthly_for_future_prediction(months: list[str]) -> pd.DataFrame:
    rows = []
    for month_index, month in enumerate(months, start=1):
        for site_no in ("8000", "8001"):
            for item_id in ("BOOK-1", "BOOK-2"):
                rows.append({
                    "month": month,
                    "site_no": site_no,
                    "blt_site_no": site_no,
                    "item_id": item_id,
                    "isbn": item_id,
                    "gds_no": item_id,
                    "gds_ctgry_3_lvel": "unknown",
                    "gds_ctgry_4_lvel": "unknown",
                    "gds_ctgry_5_lvel": "unknown",
                    "price": 20.0,
                    "total_qty": float(month_index),
                    "offline_qty": float(month_index),
                    "online_qty": 0.0,
                    "unknown_channel_qty": 0.0,
                    "total_tlp": 20.0 * month_index,
                    "total_tsp": 18.0 * month_index,
                    "avg_real_price": 18.0,
                    "discount_rate": 0.9,
                    "sales_days": 1,
                    "sales_count": 1,
                    "return_count": 0,
                    "return_qty": 0.0,
                })
    return pd.DataFrame(rows)


def test_ready_dataset_creates_e0_e2_e3_prediction_artifacts(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setenv("WENXUAN_SOFTWARE_RUNTIME", str(tmp_path / "runtime"))
    from src.models.lightgbm_model import load_model_bundle
    from software.backend.api.main import create_app
    from software.backend.services.dataset_registry import DatasetRegistry

    model_paths = [
        "models/final/two_stage_classifier_active_store_1m.txt",
        "models/checkpoints/two_stage_optimization/train_valid_protocol/selected/E0_VALID_SELECTED.txt",
        "models/checkpoints/two_stage_optimization/train_valid_protocol/selected/E2_CROSS_STORE_VALID_SELECTED.txt",
        "models/checkpoints/two_stage_optimization/train_valid_protocol/selected/E3_DIFF_CROSS_STORE_VALID_SELECTED.txt",
    ]
    features: set[str] = set()
    for path in model_paths:
        _, metadata = load_model_bundle(path)
        features.update(metadata["feature_names"])
    row = {column: 0.0 for column in features}
    row.update({
        "month": "2026-03", "site_no": "8000", "item_id": "BOOK-1", "future_qty_1m": 0.0,
        "future_qty_2m": 0.0, "blt_site_no": "8000", "gds_ctgry_3_lvel": "unknown",
        "gds_ctgry_4_lvel": "unknown", "gds_ctgry_5_lvel": "unknown", "book_name": "Smoke Book",
    })
    feature_path = tmp_path / "features.parquet"
    pd.DataFrame([row]).to_parquet(feature_path, index=False)
    dataset = DatasetRegistry().register({
        "dataset_name": "prediction-smoke", "source_type": "csv", "source_files": [],
        "date_range": {"start": "2026-03", "end": "2026-03"}, "store_count": 1, "item_count": 1, "row_count": 1,
        "monthly_parquet_path": str(feature_path), "active_store_parquet_path": str(feature_path), "feature_parquet_path": str(feature_path),
        "has_active_store": True, "has_diff_features": True, "has_cross_store_features": True,
    })
    client = TestClient(create_app())

    created = client.post("/api/predictions", json={"dataset_id": dataset["dataset_id"], "model_ids": ["E0", "E2", "E3"]})
    assert created.status_code == 201, created.text
    payload = created.json()
    assert client.get(f"/api/jobs/{payload['job_id']}").json()["status"] == "SUCCESS"
    run = client.get(f"/api/predictions/{payload['prediction_run_id']}").json()
    assert run["status"] == "SUCCESS"
    assert Path(run["prediction_dir"], "predictions.parquet").is_file()
    summary = client.get(f"/api/predictions/{payload['prediction_run_id']}/summary").json()
    assert set(summary["model_summaries"]) == {"E0", "E2", "E3"}
    assert summary["observation_month"] == "2026-03"
    assert summary["target_month"] == "2026-04"
    filtered_summary = client.get(
        f"/api/predictions/{payload['prediction_run_id']}/summary?model_id=E2&site_no=8000&mc=MC0"
    )
    assert filtered_summary.status_code == 200, filtered_summary.text
    assert filtered_summary.json()["filtered_summary"]["filters"] == {"model_id": "E2", "site_no": "8000", "mc": "MC0"}
    assert filtered_summary.json()["filtered_summary"]["store_count"] == 1
    assert set(filtered_summary.json()["filtered_summary"]["mc_counts"]) == {"MC0", "MC1", "MC2", "MC3", "MC4"}
    top_books = client.get(f"/api/predictions/{payload['prediction_run_id']}/results?kind=top_books").json()
    assert top_books["total"] == 3
    assert {row["model_id"] for row in top_books["items"]} == {"E0", "E2", "E3"}
    downloaded = client.get(f"/api/predictions/{payload['prediction_run_id']}/download-parquet")
    assert downloaded.status_code == 200

    filtered = client.get(
        f"/api/predictions/{payload['prediction_run_id']}/results?kind=predictions&model_id=E2&site_no=8000&mc=MC0&page=1&page_size=10&sort_by=pred_qty_int&sort_order=desc"
    )
    assert filtered.status_code == 200, filtered.text
    assert filtered.json()["total"] == 1
    assert filtered.json()["items"][0]["model_id"] == "E2"
    assert filtered.json()["sort"] == {"sort_by": "pred_qty_int", "sort_order": "desc"}

    store_summary = client.get(f"/api/predictions/{payload['prediction_run_id']}/store/8000/summary?model_id=E2")
    assert store_summary.status_code == 200, store_summary.text
    assert store_summary.json()["prediction_total"] >= 0

    exported = client.get(f"/api/predictions/{payload['prediction_run_id']}/export-excel?model_id=E2&site_no=8000&mc=MC0&include_predictions=true")
    assert exported.status_code == 200, exported.text
    assert exported.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert f"prediction_{payload['prediction_run_id']}_E2_site8000_MC0.xlsx" in exported.headers["content-disposition"]
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    assert {"Summary", "Store_Summary", "Top_Books", "Predictions", "Model_Info"}.issubset(workbook.sheetnames)
    assert workbook["Predictions"].max_row == 2

    difficult = client.get(f"/api/predictions/{payload['prediction_run_id']}/difficult-books?model_id=E2")
    assert difficult.status_code == 200, difficult.text
    difficult_payload = difficult.json()
    assert difficult_payload["total"] >= 0

    difficult_summary = client.get(f"/api/predictions/{payload['prediction_run_id']}/difficult-books/summary?model_id=E2")
    assert difficult_summary.status_code == 200, difficult_summary.text
    assert difficult_summary.json()["difficulty_rules"]["metric"] == "stable_error"


def test_future_prediction_uses_standard_history_without_future_labels(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setenv("WENXUAN_SOFTWARE_RUNTIME", str(tmp_path / "runtime"))
    from software.backend.api.main import create_app
    from software.backend.services.dataset_registry import DatasetRegistry

    monthly_path = tmp_path / "monthly.parquet"
    _monthly_for_future_prediction(["2025-12", "2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06"]).to_parquet(monthly_path, index=False)
    DatasetRegistry().register({
        "dataset_name": "standard-history-smoke",
        "source_type": "parquet",
        "source_files": ["raw-file-1"],
        "date_range": {"start": "2025-12", "end": "2026-06"},
        "store_count": 2,
        "item_count": 2,
        "row_count": 28,
        "monthly_parquet_path": str(monthly_path),
        "active_store_parquet_path": str(monthly_path),
        "feature_parquet_path": str(monthly_path),
        "has_active_store": True,
        "has_diff_features": True,
        "has_cross_store_features": True,
    })
    client = TestClient(create_app())

    created = client.post("/api/predictions", json={"target_month": "2026-07", "model_ids": ["E0", "E2", "E3"]})

    assert created.status_code == 201, created.text
    payload = created.json()
    summary = client.get(f"/api/predictions/{payload['prediction_run_id']}/summary").json()
    assert summary["observation_month"] == "2026-06"
    assert summary["target_month"] == "2026-07"
    assert summary["provenance"]["source_months"] == ["2025-12", "2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06"]
    run = client.get(f"/api/predictions/{payload['prediction_run_id']}").json()
    inference_path = Path(run["inference_feature_path"])
    assert inference_path.is_file()
    inference_columns = set(pd.read_parquet(inference_path).columns)
    assert "future_qty_1m" not in inference_columns
    assert "future_qty_2m" not in inference_columns
    assert "target_qty_1m" not in inference_columns


def test_difficult_books_endpoint_handles_actual_zero_and_exports(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setenv("WENXUAN_SOFTWARE_RUNTIME", str(tmp_path / "runtime"))
    from software.backend.api.main import create_app
    from software.backend.services.dataset_registry import DatasetRegistry
    from software.backend.services.prediction_registry import PredictionRegistry

    prediction_dir = tmp_path / "runtime" / "predictions" / "manual-run"
    prediction_dir.mkdir(parents=True)
    monthly_path = tmp_path / "monthly.parquet"
    pd.DataFrame([
        {"month": "2026-01", "site_no": "8000", "item_id": "BOOK-ZERO", "total_qty": 3},
        {"month": "2026-02", "site_no": "8000", "item_id": "BOOK-ZERO", "total_qty": 5},
    ]).to_parquet(monthly_path, index=False)
    DatasetRegistry().register({
        "dataset_id": "dataset-1",
        "dataset_name": "manual-dataset",
        "source_type": "csv",
        "source_files": [],
        "date_range": {"start": "2026-01", "end": "2026-02"},
        "store_count": 1,
        "item_count": 1,
        "row_count": 2,
        "monthly_parquet_path": str(monthly_path),
        "active_store_parquet_path": str(monthly_path),
        "feature_parquet_path": str(monthly_path),
        "has_active_store": True,
        "has_diff_features": True,
        "has_cross_store_features": True,
    })
    pd.DataFrame([
        {
            "dataset_id": "dataset-1",
            "prediction_run_id": "manual-run",
            "model_id": "E2",
            "site_no": "8000",
            "item_id": "BOOK-ZERO",
            "book_name": "Zero Actual Book",
            "p_sale": 0.9,
            "conditional_qty": 12.0,
            "pred_qty_raw": 12.0,
            "pred_qty_int": 12,
            "pred_mc": "MC3",
            "actual_qty": 0,
        }
    ]).to_parquet(prediction_dir / "predictions.parquet", index=False)
    pd.DataFrame().to_parquet(prediction_dir / "top_books.parquet", index=False)
    pd.DataFrame().to_parquet(prediction_dir / "store_summary.parquet", index=False)
    (prediction_dir / "summary.json").write_text(
        '{"prediction_run_id":"manual-run","dataset_id":"dataset-1","observation_month":"2026-03","model_ids":["E2"],"model_summaries":{}}',
        encoding="utf-8",
    )
    registry = PredictionRegistry()
    registry.create({
        "prediction_run_id": "manual-run",
        "dataset_id": "dataset-1",
        "model_ids": ["E2"],
        "observation_month": "2026-03",
        "store_ids": None,
        "job_id": "job-1",
    })
    registry.update("manual-run", status="SUCCESS", prediction_dir=str(prediction_dir))
    client = TestClient(create_app())

    difficult = client.get("/api/predictions/manual-run/difficult-books?model_id=E2&page=1&page_size=10")
    assert difficult.status_code == 200, difficult.text
    payload = difficult.json()
    assert payload["total"] == 1
    assert payload["items"][0]["actual_qty"] == 0
    assert payload["items"][0]["stable_error"] == 12
    assert payload["items"][0]["difficulty_level"] == "困难"

    summary = client.get("/api/predictions/manual-run/difficult-books/summary?model_id=E2")
    assert summary.status_code == 200, summary.text
    assert summary.json()["difficult_book_count"] == 1
    assert summary.json()["difficulty_rules"]["note"].startswith("当前为系统工程默认筛选规则")

    difficult_export = client.get("/api/predictions/manual-run/difficult-books/export?model_id=E2")
    assert difficult_export.status_code == 200, difficult_export.text
    assert difficult_export.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    series = client.get("/api/predictions/manual-run/historical-series?model_id=E2&site_no=8000&item_id=BOOK-ZERO")
    assert series.status_code == 200, series.text
    assert series.json()["items"] == [
        {"month": "2026-01", "actual_qty": 3, "pred_qty": None, "is_prediction": False},
        {"month": "2026-02", "actual_qty": 5, "pred_qty": None, "is_prediction": False},
        {"month": "2026-04", "actual_qty": None, "pred_qty": 12, "is_prediction": True},
    ]
